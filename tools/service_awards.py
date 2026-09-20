import re
from calendar import monthrange
from datetime import date, datetime

from openpyxl import load_workbook


INACTIVE_STATUSES = {
    "terminated",
    "deceased",
    "resigned",
    "retired",
    "inactive",
}

FULL_TIME_VALUES = {
    "ft",
    "full time",
    "full-time",
    "regular full time",
    "regular full-time",
}

NON_FULL_TIME_VALUES = {
    "pt",
    "part time",
    "part-time",
    "prn",
    "per diem",
    "temporary",
}


def _clean_header(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def _clean_value(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).casefold()


def _as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value in (None, ""):
        return None
    text = str(value).strip()
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _column(columns, aliases, label, required=False):
    for alias in aliases:
        if alias in columns:
            return columns[alias]
    if required:
        raise ValueError(f"UKG report is missing the {label} column.")
    return None


def _find_header_row(rows):
    for index, row in enumerate(rows[:30]):
        headers = {_clean_header(value) for value in row}
        has_name = (
            {"firstname", "lastname"}.issubset(headers)
            or "employeename" in headers
            or "name" in headers
        )
        has_date = bool(headers.intersection({
            "latestdate",
            "latesthireorrehiredate",
            "datehired",
            "hiredate",
            "adjustedbenefitdate",
            "benefitdate",
            "servicedate",
            "senioritydate",
        }))
        if has_name and has_date:
            return index
    raise ValueError(
        "Could not find the employee headings. Include employee name and a hire, service, or benefit date."
    )


def parse_milestones(value):
    pieces = re.findall(r"\d+", str(value or ""))
    milestones = sorted({int(piece) for piece in pieces if int(piece) > 0})
    if not milestones:
        raise ValueError("Enter at least one service-award milestone, such as 5, 10, 15.")
    return milestones


def _anniversary(service_date, year):
    day = min(service_date.day, monthrange(year, service_date.month)[1])
    return date(year, service_date.month, day)


def _employment_classification(value):
    normalized = _clean_value(value)
    if normalized in FULL_TIME_VALUES or "full time" in normalized or "full-time" in normalized:
        return "full_time"
    if normalized in NON_FULL_TIME_VALUES or normalized.startswith("prn") or "part time" in normalized:
        return "not_full_time"
    return "unknown"


def read_service_award_employees(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        header_row = _find_header_row(rows)
        columns = {
            _clean_header(value): index
            for index, value in enumerate(rows[header_row])
            if value not in (None, "")
        }

        first_col = _column(columns, {"firstname", "first"}, "First Name")
        last_col = _column(columns, {"lastname", "last"}, "Last Name")
        name_col = _column(columns, {"employeename", "name"}, "Employee Name")
        if name_col is None and (first_col is None or last_col is None):
            raise ValueError("UKG report must include First Name and Last Name, or Employee Name.")

        employee_id_col = _column(
            columns,
            {"employeeid", "employeenumber", "employeeno", "empid", "id"},
            "Employee Number",
        )
        status_col = _column(columns, {"employeestatus", "status"}, "Employee Status")
        employment_type_col = _column(
            columns,
            {
                "employeetype",
                "employmenttype",
                "fullparttime",
                "fulltimeparttime",
                "regulartemporary",
            },
            "Employee Type",
        )
        gender_col = _column(columns, {"gender", "sex"}, "Gender")

        date_choices = [
            ("Adjusted Benefit Date", _column(columns, {"adjustedbenefitdate"}, "Adjusted Benefit Date")),
            ("Service Date", _column(columns, {"servicedate"}, "Service Date")),
            ("Seniority Date", _column(columns, {"senioritydate"}, "Seniority Date")),
            ("Benefit Date", _column(columns, {"benefitdate", "benefitstartdate"}, "Benefit Date")),
            (
                "Latest Hire/Rehire Date",
                _column(columns, {"latestdate", "latesthireorrehiredate"}, "Latest Date"),
            ),
            ("Rehire Date", _column(columns, {"daterehired", "rehiredate"}, "Rehire Date")),
            ("Hire Date", _column(columns, {"datehired", "hiredate"}, "Hire Date")),
        ]
        if all(column is None for _, column in date_choices):
            raise ValueError("UKG report is missing a hire, service, or benefit date.")

        employees = []
        for row in rows[header_row + 1:]:
            if name_col is not None:
                name = str(row[name_col] or "").strip()
            else:
                first = str(row[first_col] or "").strip()
                last = str(row[last_col] or "").strip()
                name = f"{first.title()} {last.title()}".strip()
            if not name:
                continue

            service_date = None
            date_source = ""
            for label, column in date_choices:
                if column is None:
                    continue
                candidate = _as_date(row[column])
                if candidate:
                    service_date = candidate
                    date_source = label
                    break

            employees.append({
                "name": name,
                "employee_id": str(row[employee_id_col] or "").strip() if employee_id_col is not None else "",
                "status": str(row[status_col] or "").strip() if status_col is not None else "",
                "employment_type": (
                    str(row[employment_type_col] or "").strip()
                    if employment_type_col is not None
                    else ""
                ),
                "gender": str(row[gender_col] or "").strip() if gender_col is not None else "",
                "service_date": service_date,
                "date_source": date_source,
            })
        if not employees:
            raise ValueError("No employee records were found in the UKG report.")
        return employees
    finally:
        workbook.close()


def service_award_review(employees, award_year, milestones):
    eligible = []
    needs_review = []
    excluded = {
        "inactive": 0,
        "not_full_time": 0,
        "not_milestone": 0,
        "missing_date": 0,
    }

    for employee in employees:
        status = _clean_value(employee["status"])
        if status in INACTIVE_STATUSES:
            excluded["inactive"] += 1
            continue

        service_date = employee["service_date"]
        if service_date is None:
            item = {**employee, "review_reason": "Missing service/hire date"}
            needs_review.append(item)
            excluded["missing_date"] += 1
            continue

        years = award_year - service_date.year
        if years not in milestones or years <= 0:
            excluded["not_milestone"] += 1
            continue

        item = {
            **employee,
            "years": years,
            "anniversary": _anniversary(service_date, award_year),
        }
        classification = _employment_classification(employee["employment_type"])
        if classification == "not_full_time":
            excluded["not_full_time"] += 1
            continue
        if classification == "unknown":
            item["review_reason"] = "Verify full-time status"
            needs_review.append(item)
            continue
        if not employee["employee_id"]:
            item["review_reason"] = "Missing employee number"
            needs_review.append(item)
            continue
        if not employee["gender"]:
            item["review_reason"] = "Missing gender"
            needs_review.append(item)
            continue
        eligible.append(item)

    sort_key = lambda item: (
        item.get("years", 0),
        item.get("anniversary") or date.max,
        item["name"].casefold(),
    )
    eligible.sort(key=sort_key)
    needs_review.sort(key=sort_key)
    return eligible, needs_review, excluded
