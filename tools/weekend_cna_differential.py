from collections import defaultdict
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import load_workbook


WEEKEND_RATE = Decimal("4.00")


def _normalize(value):
    if value is None:
        return ""
    return str(value).strip().replace("\n", " ")


def _as_number(value):
    if isinstance(value, (int, float)):
        return float(value)
    if value in (None, "", "-"):
        return None
    try:
        return float(str(value).replace(",", "").strip())
    except (TypeError, ValueError):
        return None


def _as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        value = value.strip()
        for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d"):
            try:
                return datetime.strptime(value, fmt).date()
            except ValueError:
                pass
    return None


def calculate_weekend_cna_differential(xlsx_path, rate=WEEKEND_RATE):
    """Return weekend CNA hours and $4/hr differential totals from a UKG report."""
    workbook = load_workbook(xlsx_path, data_only=True, read_only=True)
    sheet = workbook.active

    header_row = None
    headers = {}
    for row_num, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        normalized = [_normalize(value) for value in row]
        if "First Name" in normalized and "Last Name" in normalized and "Total Work Hours" in normalized:
            header_row = row_num
            headers = {name: idx for idx, name in enumerate(normalized) if name}
            break

    if header_row is None:
        raise ValueError(
            "I could not find the UKG employee-hours header row. "
            "Use the Calculated Hours By Work Day report that includes First Name, Last Name, Date, and Total Work Hours."
        )

    required = ["First Name", "Last Name", "Date", "Total Work Hours"]
    missing = [column for column in required if column not in headers]
    if missing:
        raise ValueError("The report is missing required column(s): " + ", ".join(missing))

    totals = defaultdict(lambda: {"hours": 0.0, "dates": set()})
    skipped_non_cna = 0
    skipped_non_weekend = 0

    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        first = _normalize(row[headers["First Name"]]) if headers["First Name"] < len(row) else ""
        last = _normalize(row[headers["Last Name"]]) if headers["Last Name"] < len(row) else ""
        if not first and not last:
            continue

        if "Location(5)" in headers and headers["Location(5)"] < len(row):
            location5 = _normalize(row[headers["Location(5)"]]).upper()
            if location5 and location5 != "CNA":
                skipped_non_cna += 1
                continue

        work_date = _as_date(row[headers["Date"]]) if headers["Date"] < len(row) else None
        if work_date is None:
            continue
        if work_date.weekday() not in (5, 6):  # Saturday / Sunday
            skipped_non_weekend += 1
            continue

        hours = _as_number(row[headers["Total Work Hours"]]) if headers["Total Work Hours"] < len(row) else None
        if hours is None or hours <= 0:
            continue

        employee_id = ""
        if "Employee Id" in headers and headers["Employee Id"] < len(row):
            employee_id = _normalize(row[headers["Employee Id"]])

        display_name = " ".join(part for part in (first.title(), last.title()) if part)
        key = employee_id or display_name.upper()
        totals[key]["name"] = display_name
        totals[key]["employee_id"] = employee_id
        totals[key]["hours"] += hours
        totals[key]["dates"].add(work_date)

    rate_decimal = Decimal(str(rate))
    results = []
    total_hours = 0.0
    total_owed = Decimal("0.00")

    for item in totals.values():
        hours = round(item["hours"], 2)
        owed = (Decimal(str(hours)) * rate_decimal).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        results.append(
            {
                "employee_id": item["employee_id"],
                "name": item["name"],
                "hours": hours,
                "owed": float(owed),
                "dates": sorted(item["dates"]),
            }
        )
        total_hours += hours
        total_owed += owed

    results.sort(key=lambda item: item["name"].lower())

    return {
        "rate": float(rate_decimal),
        "results": results,
        "total_hours": round(total_hours, 2),
        "total_owed": float(total_owed.quantize(Decimal("0.01"))),
        "employee_count": len(results),
        "skipped_non_cna": skipped_non_cna,
        "skipped_non_weekend": skipped_non_weekend,
    }
